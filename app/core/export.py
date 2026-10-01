"""Convert the LLM's markdown resume into downloadable Word (.docx) or PDF files."""
import re
from io import BytesIO

from docx import Document
from docx.shared import Pt
from fpdf import FPDF

FORMATS = {
    "Word (.docx)": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "PDF (.pdf)": ("pdf", "application/pdf"),
    "Markdown (.md)": ("md", "text/markdown"),
}


def _lines(markdown: str):
    """Yield (kind, text) for each meaningful markdown line."""
    for raw in markdown.splitlines():
        line = raw.strip()
        if not line or line == "---":
            continue
        if line.startswith("### "):
            yield "h3", line[4:].strip()
        elif line.startswith("## "):
            yield "h2", line[3:].strip()
        elif line.startswith("# "):
            yield "h1", line[2:].strip()
        elif line[:2] in ("- ", "* ", "• "):
            yield "bullet", line[2:].strip()
        else:
            yield "text", line


# ---------- Word ----------

def _add_runs(paragraph, text: str) -> None:
    """Add text to a paragraph, turning **bold** markdown into real bold."""
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def markdown_to_docx(markdown: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for kind, text in _lines(markdown):
        if kind == "h1":
            doc.add_heading(text, level=0)
        elif kind == "h2":
            doc.add_heading(text, level=1)
        elif kind == "h3":
            doc.add_heading(text, level=2)
        elif kind == "bullet":
            _add_runs(doc.add_paragraph(style="List Bullet"), text)
        else:
            _add_runs(doc.add_paragraph(), text)

    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------- PDF ----------

_REPLACEMENTS = {
    "‘": "'", "’": "'", "“": '"', "”": '"',
    "–": "-", "—": "-", "•": "-", "…": "...", " ": " ",
}


def _pdf_safe(text: str) -> str:
    """The built-in PDF fonts support Latin-1 (English, French, Spanish, ...)."""
    for bad, good in _REPLACEMENTS.items():
        text = text.replace(bad, good)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def markdown_to_pdf(markdown: str) -> bytes:
    pdf = FPDF(format="A4")
    pdf.set_margins(18, 15, 18)
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    for kind, text in _lines(markdown):
        text = _pdf_safe(text)
        pdf.set_x(pdf.l_margin)
        if kind == "h1":
            pdf.set_font("Helvetica", "B", 20)
            pdf.multi_cell(width, 10, text.replace("**", ""), align="C")
        elif kind == "h2":
            pdf.ln(3)
            pdf.set_font("Helvetica", "B", 12)
            pdf.set_text_color(30, 64, 175)
            pdf.multi_cell(width, 7, text.replace("**", "").upper())
            pdf.set_text_color(0, 0, 0)
            y = pdf.get_y()
            pdf.set_draw_color(30, 64, 175)
            pdf.line(pdf.l_margin, y, pdf.l_margin + width, y)
            pdf.ln(2)
        elif kind == "h3":
            pdf.ln(1)
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(width, 6, text.replace("**", ""))
        elif kind == "bullet":
            pdf.set_font("Helvetica", "", 10.5)
            pdf.cell(5, 5.5, "-")
            pdf.multi_cell(width - 5, 5.5, text, markdown=True)
        else:
            pdf.set_font("Helvetica", "", 10.5)
            align = "C" if pdf.page_no() == 1 and pdf.get_y() < 40 else "L"  # contact line under the name
            pdf.multi_cell(width, 5.5, text, markdown=True, align=align)

    return bytes(pdf.output())


def export(markdown: str, fmt: str) -> bytes:
    """Return the resume in the chosen format: 'docx', 'pdf' or 'md'."""
    if fmt == "docx":
        return markdown_to_docx(markdown)
    if fmt == "pdf":
        return markdown_to_pdf(markdown)
    return markdown.encode("utf-8")
